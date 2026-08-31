from __future__ import annotations

import hashlib
import os
import re
import stat
import tempfile
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any, Callable, ContextManager
from xml.etree import ElementTree


TEXT_MAX_BYTES = 5 * 1024 * 1024
DOCX_MAX_BYTES = 20 * 1024 * 1024
PDF_MAX_BYTES = 50 * 1024 * 1024
MAX_DOCUMENT_UPLOAD_BYTES = PDF_MAX_BYTES
MAX_DOCX_ENTRIES = 512
MAX_DOCX_EXPANSION_RATIO = 100
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_NORMALIZED_TEXT_BYTES = 10 * 1024 * 1024

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
FORMAT_SPECS = {
    ".txt": ("text/plain", TEXT_MAX_BYTES, "utf-8"),
    ".md": ("text/markdown", TEXT_MAX_BYTES, "utf-8"),
    ".pdf": ("application/pdf", PDF_MAX_BYTES, "none"),
    ".docx": (DOCX_CONTENT_TYPE, DOCX_MAX_BYTES, "docx"),
}

_PACKAGE_RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
_OFFICE_DOCUMENT_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
)
_DOCX_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)
_SAFE_DIGEST = re.compile(r"^[0-9a-f]{64}$")


class DocumentValidationError(ValueError):
    pass


class DocumentTooLargeError(DocumentValidationError):
    pass


@dataclass(frozen=True)
class ValidatedDocument:
    original_filename: str
    content_type: str
    extension: str
    size: int
    sha256: str
    extraction_kind: str
    extraction_status: str
    normalized_text: str | None
    normalized_text_digest: str | None


def validate_document(
    filename: str,
    content_type: str,
    content: bytes,
) -> ValidatedDocument:
    original_filename, normalized_content_type, extension, spec = _document_spec(
        filename, content_type
    )
    _expected_content_type, max_bytes, extraction_kind = spec
    if not content:
        raise DocumentValidationError("Document must not be empty")
    if len(content) > max_bytes:
        raise DocumentTooLargeError("Document exceeds the format size limit")

    normalized_text: str | None
    extraction_status: str
    if extraction_kind == "utf-8":
        normalized_text = _decode_text(content)
        extraction_status = "complete"
    elif extraction_kind == "docx":
        _validate_docx_container(content)
        normalized_text = _extract_docx_text(content)
        extraction_status = "complete"
    else:
        if not content.startswith(b"%PDF-"):
            raise DocumentValidationError("PDF bytes do not match the declared type")
        normalized_text = None
        extraction_status = "not-applicable"

    normalized_text_digest = None
    if normalized_text is not None:
        normalized_bytes = normalized_text.encode("utf-8")
        if len(normalized_bytes) > MAX_NORMALIZED_TEXT_BYTES:
            raise DocumentValidationError("Extracted document text exceeds the safety limit")
        normalized_text_digest = hashlib.sha256(normalized_bytes).hexdigest()
    return ValidatedDocument(
        original_filename=original_filename,
        content_type=normalized_content_type,
        extension=extension,
        size=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        extraction_kind=extraction_kind,
        extraction_status=extraction_status,
        normalized_text=normalized_text,
        normalized_text_digest=normalized_text_digest,
    )


def document_size_limit(filename: str, content_type: str) -> int:
    _filename, _content_type, _extension, spec = _document_spec(filename, content_type)
    return spec[1]


def _document_spec(
    filename: str,
    content_type: str,
) -> tuple[str, str, str, tuple[str, int, str]]:
    original_filename = _validate_filename(filename)
    extension = Path(original_filename).suffix.lower()
    spec = FORMAT_SPECS.get(extension)
    if spec is None:
        raise DocumentValidationError("Unsupported document extension")
    expected_content_type = spec[0]
    normalized_content_type = content_type.split(";", 1)[0].strip().lower()
    if normalized_content_type != expected_content_type:
        raise DocumentValidationError("Document extension and content type do not match")
    return original_filename, normalized_content_type, extension, spec


def _validate_filename(filename: str) -> str:
    if not isinstance(filename, str):
        raise DocumentValidationError("Document filename is invalid")
    normalized = unicodedata.normalize("NFC", filename.strip())
    if (
        not normalized
        or len(normalized) > 255
        or normalized in {".", ".."}
        or "\x00" in normalized
        or "/" in normalized
        or "\\" in normalized
        or Path(normalized).name != normalized
    ):
        raise DocumentValidationError("Document filename is invalid")
    return normalized


def _decode_text(content: bytes) -> str:
    try:
        text = content.decode("utf-8", errors="strict")
    except UnicodeDecodeError as exc:
        raise DocumentValidationError("Text document must be strict UTF-8") from exc
    if "\x00" in text or any(
        ord(character) < 32 and character not in {"\t", "\n", "\r"}
        for character in text
    ):
        raise DocumentValidationError("Text document contains binary control bytes")
    return unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))


def _validate_docx_container(content: bytes) -> None:
    if not content.startswith(b"PK\x03\x04"):
        raise DocumentValidationError("DOCX bytes do not start with a ZIP local file header")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise DocumentValidationError("DOCX contains too many entries")
            names: set[str] = set()
            total_compressed = 0
            total_uncompressed = 0
            for entry in entries:
                name = entry.filename
                _validate_docx_entry(entry)
                if name in names:
                    raise DocumentValidationError("DOCX contains duplicate entries")
                names.add(name)
                total_compressed += entry.compress_size
                total_uncompressed += entry.file_size
                if entry.file_size and (
                    entry.compress_size == 0
                    or entry.file_size / entry.compress_size > MAX_DOCX_EXPANSION_RATIO
                ):
                    raise DocumentValidationError("DOCX entry expansion ratio is unsafe")
            if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentValidationError("DOCX uncompressed content exceeds the safety limit")
            if total_uncompressed and (
                total_compressed == 0
                or total_uncompressed / total_compressed > MAX_DOCX_EXPANSION_RATIO
            ):
                raise DocumentValidationError("DOCX expansion ratio is unsafe")
            required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
            if not required <= names:
                raise DocumentValidationError("DOCX package is missing required Office parts")
            _validate_content_types(archive.read("[Content_Types].xml"))
            _validate_root_relationships(archive.read("_rels/.rels"))
            for name in names:
                if name.endswith(".rels"):
                    _reject_external_relationships(archive.read(name))
            ElementTree.fromstring(archive.read("word/document.xml"))
    except DocumentValidationError:
        raise
    except (OSError, KeyError, RuntimeError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise DocumentValidationError("DOCX container is corrupt or encrypted") from exc


def _validate_docx_entry(entry: zipfile.ZipInfo) -> None:
    name = entry.filename
    path = PurePosixPath(name)
    if (
        not name
        or "\x00" in name
        or "\\" in name
        or name.startswith("/")
        or path.is_absolute()
        or any(part in {"", ".", ".."} for part in path.parts)
        or (path.parts and ":" in path.parts[0])
    ):
        raise DocumentValidationError("DOCX contains an unsafe entry path")
    if entry.flag_bits & 0x1:
        raise DocumentValidationError("Encrypted DOCX files are not supported")
    unix_mode = (entry.external_attr >> 16) & 0xFFFF
    if unix_mode and stat.S_IFMT(unix_mode) not in {0, stat.S_IFREG, stat.S_IFDIR}:
        raise DocumentValidationError("DOCX contains a special-file entry")


def _validate_content_types(content: bytes) -> None:
    root = ElementTree.fromstring(content)
    overrides = {
        element.attrib.get("PartName"): element.attrib.get("ContentType")
        for element in root.findall(f"{{{_CONTENT_TYPES}}}Override")
    }
    if overrides.get("/word/document.xml") != _DOCX_MAIN_CONTENT_TYPE:
        raise DocumentValidationError("DOCX main Office content type is invalid")


def _validate_root_relationships(content: bytes) -> None:
    root = ElementTree.fromstring(content)
    office_relationships = [
        relationship
        for relationship in root.findall(f"{{{_PACKAGE_RELATIONSHIPS}}}Relationship")
        if relationship.attrib.get("Type") == _OFFICE_DOCUMENT_RELATIONSHIP
    ]
    if len(office_relationships) != 1:
        raise DocumentValidationError("DOCX Office document relationship is invalid")
    relationship = office_relationships[0]
    if (
        relationship.attrib.get("TargetMode", "Internal").lower() == "external"
        or PurePosixPath(relationship.attrib.get("Target", "")).as_posix()
        != "word/document.xml"
    ):
        raise DocumentValidationError("DOCX Office document relationship is unsafe")


def _reject_external_relationships(content: bytes) -> None:
    root = ElementTree.fromstring(content)
    for relationship in root.findall(f"{{{_PACKAGE_RELATIONSHIPS}}}Relationship"):
        if relationship.attrib.get("TargetMode", "Internal").lower() == "external":
            raise DocumentValidationError("DOCX external relationships are not supported")


def _extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ModuleNotFoundError as exc:
        raise DocumentValidationError("DOCX decoder is unavailable") from exc
    try:
        document = Document(BytesIO(content))
        lines: list[str] = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                if block.text:
                    lines.append(block.text)
            elif isinstance(block, Table):
                for row in block.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
        return unicodedata.normalize("NFC", "\n".join(lines))
    except (KeyError, ValueError, OSError, zipfile.BadZipFile) as exc:
        raise DocumentValidationError("DOCX text extraction failed") from exc


class DocumentResourceStore:
    def __init__(
        self,
        data_dir: Path,
        connect: Callable[[], ContextManager[Any]],
        transaction: Callable[[], ContextManager[Any]],
        require_active_project: Callable[[Any, str], None],
        now_ms: Callable[[], int],
        consistency_lock: ContextManager[Any],
    ) -> None:
        self.data_dir = data_dir
        self.documents_dir = data_dir / "documents"
        self._connect = connect
        self._transaction = transaction
        self._require_active_project = require_active_project
        self._now_ms = now_ms
        self._consistency_lock = consistency_lock

    def create(
        self,
        project_id: str,
        filename: str,
        content_type: str,
        content: bytes,
    ) -> dict[str, Any]:
        validated = validate_document(filename, content_type, content)
        with self._consistency_lock:
            return self._create_validated(project_id, content, validated)

    def _create_validated(
        self,
        project_id: str,
        content: bytes,
        validated: ValidatedDocument,
    ) -> dict[str, Any]:
        self.documents_dir.mkdir(parents=True, exist_ok=True)
        resource_id = uuid.uuid4().hex
        stored_name = f"{resource_id}{validated.extension}"
        relative_path = f"documents/{stored_name}"
        final_path = self.documents_dir / stored_name
        fd, temp_name = tempfile.mkstemp(
            prefix=".document-", suffix=".tmp", dir=str(self.documents_dir)
        )
        now = self._now_ms()
        try:
            with os.fdopen(fd, "wb") as handle:
                handle.write(content)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temp_name, final_path)
            with self._transaction() as connection:
                self._require_active_project(connection, project_id)
                connection.execute(
                    """INSERT INTO project_document_resources(
                           resource_id, project_id, relative_path, original_filename,
                           content_type, size, sha256, extraction_kind, extraction_status,
                           normalized_text, normalized_text_digest, revision,
                           lifecycle_status, created_at, updated_at
                       ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, 'active', ?, ?)""",
                    (
                        resource_id,
                        project_id,
                        relative_path,
                        validated.original_filename,
                        validated.content_type,
                        validated.size,
                        validated.sha256,
                        validated.extraction_kind,
                        validated.extraction_status,
                        validated.normalized_text,
                        validated.normalized_text_digest,
                        now,
                        now,
                    ),
                )
        except Exception:
            Path(temp_name).unlink(missing_ok=True)
            final_path.unlink(missing_ok=True)
            raise
        return self.get(project_id, resource_id)

    def stage_project_deletion(
        self,
        connection: Any,
        project_ids: list[str],
    ) -> tuple[Path, list[tuple[Path, Path]]] | None:
        if not project_ids:
            return None
        placeholders = ",".join("?" for _ in project_ids)
        relative_paths = [
            row[0]
            for row in connection.execute(
                f"""SELECT relative_path FROM project_document_resources
                    WHERE project_id IN ({placeholders})
                    ORDER BY resource_id""",
                tuple(project_ids),
            )
        ]
        if not relative_paths:
            return None
        staging_dir = self.data_dir / f".documents-delete-{uuid.uuid4().hex}"
        staging_dir.mkdir()
        staged: list[tuple[Path, Path]] = []
        try:
            for relative_path in relative_paths:
                source = self._safe_path(relative_path)
                temporary = staging_dir / source.name
                os.replace(source, temporary)
                staged.append((temporary, source))
        except Exception:
            self.restore_staged_deletion((staging_dir, staged))
            raise
        return staging_dir, staged

    @staticmethod
    def restore_staged_deletion(
        deletion: tuple[Path, list[tuple[Path, Path]]],
    ) -> None:
        staging_dir, staged = deletion
        for temporary, original in reversed(staged):
            if temporary.exists():
                os.replace(temporary, original)
        staging_dir.rmdir()

    @staticmethod
    def discard_staged_deletion(
        deletion: tuple[Path, list[tuple[Path, Path]]],
    ) -> None:
        staging_dir, staged = deletion
        for temporary, _original in staged:
            temporary.unlink(missing_ok=True)
        staging_dir.rmdir()

    def list(self, project_id: str) -> list[dict[str, Any]]:
        with self._connect() as connection:
            self._require_active_project(connection, project_id)
            rows = connection.execute(
                f"""SELECT {_DOCUMENT_COLUMNS} FROM project_document_resources
                    WHERE project_id=? AND lifecycle_status='active'
                    ORDER BY created_at DESC, resource_id DESC""",
                (project_id,),
            ).fetchall()
        return [_public_projection(row) for row in rows]

    def get(
        self,
        project_id: str,
        resource_id: str,
        *,
        lifecycle_status: str = "active",
    ) -> dict[str, Any]:
        with self._connect() as connection:
            self._require_active_project(connection, project_id)
            row = connection.execute(
                f"""SELECT {_DOCUMENT_COLUMNS} FROM project_document_resources
                    WHERE resource_id=? AND project_id=? AND lifecycle_status=?""",
                (resource_id, project_id, lifecycle_status),
            ).fetchone()
        if row is None:
            raise LookupError(resource_id)
        return _public_projection(row)

    def trash(self, project_id: str, resource_id: str) -> dict[str, Any]:
        return self._set_lifecycle(project_id, resource_id, "active", "trash")

    def restore(self, project_id: str, resource_id: str) -> dict[str, Any]:
        return self._set_lifecycle(project_id, resource_id, "trash", "active")

    def read(self, project_id: str, resource_id: str) -> tuple[bytes, dict[str, Any]]:
        with self._connect() as connection:
            self._require_active_project(connection, project_id)
            row = connection.execute(
                f"""SELECT {_DOCUMENT_COLUMNS} FROM project_document_resources
                    WHERE resource_id=? AND project_id=? AND lifecycle_status='active'""",
                (resource_id, project_id),
            ).fetchone()
        if row is None:
            raise LookupError(resource_id)
        path = self._safe_path(row[2])
        try:
            content = path.read_bytes()
        except OSError as exc:
            raise LookupError(resource_id) from exc
        validated = validate_document(row[3], row[4], content)
        if (
            validated.size != row[5]
            or validated.sha256 != row[6]
            or validated.normalized_text_digest != row[10]
            or not _SAFE_DIGEST.fullmatch(row[6])
        ):
            raise DocumentValidationError("Stored document integrity check failed")
        return content, {
            **_public_projection(row),
            "normalizedText": validated.normalized_text,
        }

    def _set_lifecycle(
        self,
        project_id: str,
        resource_id: str,
        current_status: str,
        next_status: str,
    ) -> dict[str, Any]:
        with self._transaction() as connection:
            self._require_active_project(connection, project_id)
            cursor = connection.execute(
                """UPDATE project_document_resources
                   SET lifecycle_status=?, revision=revision+1, updated_at=?
                   WHERE resource_id=? AND project_id=? AND lifecycle_status=?""",
                (
                    next_status,
                    self._now_ms(),
                    resource_id,
                    project_id,
                    current_status,
                ),
            )
            if cursor.rowcount != 1:
                raise LookupError(resource_id)
        return self.get(project_id, resource_id, lifecycle_status=next_status)

    def _safe_path(self, relative_path: str) -> Path:
        parts = PurePosixPath(relative_path).parts
        if len(parts) != 2 or parts[0] != "documents" or parts[1] in {"", ".", ".."}:
            raise DocumentValidationError("Stored document path is invalid")
        candidate = self.documents_dir / parts[1]
        try:
            metadata = os.lstat(candidate)
        except OSError as exc:
            raise LookupError(relative_path) from exc
        reparse_flag = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)
        attributes = getattr(metadata, "st_file_attributes", 0)
        if (
            stat.S_ISLNK(metadata.st_mode)
            or bool(attributes & reparse_flag)
            or not stat.S_ISREG(metadata.st_mode)
        ):
            raise DocumentValidationError("Stored document is not a regular file")
        return candidate


_DOCUMENT_COLUMNS = """resource_id, project_id, relative_path, original_filename,
content_type, size, sha256, extraction_kind, extraction_status, normalized_text,
normalized_text_digest, revision, lifecycle_status, created_at, updated_at"""


def _public_projection(row: Any) -> dict[str, Any]:
    return {
        "id": row[0],
        "projectId": row[1],
        "originalFilename": row[3],
        "contentType": row[4],
        "size": row[5],
        "sha256": row[6],
        "extractionKind": row[7],
        "extractionStatus": row[8],
        "normalizedTextDigest": row[10],
        "revision": row[11],
        "lifecycleStatus": row[12],
        "createdAt": row[13],
        "updatedAt": row[14],
    }
